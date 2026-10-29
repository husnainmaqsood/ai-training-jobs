"""
Excel/CSV job import, run from Django Admin (Jobs → "Import from Excel").

Flow:
1. `read_rows()` turns the upload into plain dicts (canonical column → text). Header spellings are flexible.
2. The rows are stored on a pending `ImportRun`, and `analyse()` dry-runs them against the database: every row is
   validated and normalised, matched to an existing job (or not), and classified as new / update / unchanged /
   duplicate / rejected. Nothing is written. The admin sees this as a preview.
3. On confirmation `apply()` analyses again (so the result reflects the database at that moment) and writes every
   valid row inside one transaction: if anything unexpected fails, nothing is changed.

Matching an existing job (in order):
* `external_id` + platform, when the sheet has a stable job ID;
* otherwise platform + canonical job link (see `jobs.models.canonical_url`: host case, www, scheme, trailing slash,
  fragments and tracking parameters are ignored).
A link that already belongs to a job on a *different* platform, or that matches more than one job, is rejected
rather than guessed. Matched jobs are updated in place, so their database IDs (and URLs) never change.

Empty cells never erase data on an existing job; only filled-in cells are applied. Jobs missing from a sheet are
left alone in "add & update" mode; "full snapshot" mode closes them, but only after explicit confirmation.

Nothing is ever fetched from the job links.
"""

import csv
import io
import re
from dataclasses import dataclass, field
from datetime import date, datetime, time
from decimal import Decimal, InvalidOperation

from django.db import transaction
from django.utils import timezone
from django.utils.dateparse import parse_date
from django.utils.text import slugify

from core.countries import COUNTRIES

from .categories import infer_category
from .importers import COUNTRY_NAME_TO_CODE, WORLDWIDE_WORDS, html_to_text, infer_location, safe_url
from .models import Category, ImportRun, Job, Platform, canonical_url

MAX_BYTES = 10 * 1024 * 1024
MAX_ROWS = 5000

# canonical column -> accepted header spellings (compared lower-cased; spaces, underscores and dashes ignored)
COLUMNS = {
    "external_id": ["external id", "job id", "id", "reference", "ref id", "listing id"],
    "title": ["title", "job title", "position", "role"],
    "platform": ["platform", "company", "company name", "employer"],
    "job_link": ["job link", "link", "url", "job url", "application url", "application link", "apply link", "apply url"],
    "job_description": ["job description", "description", "details", "summary"],
    "hourly_rate": ["hourly rate", "rate", "pay", "pay rate", "compensation"],
    "minimum_rate": ["minimum rate", "min rate", "rate min", "minimum pay", "min pay"],
    "maximum_rate": ["maximum rate", "max rate", "rate max", "maximum pay", "max pay"],
    "currency": ["currency"],
    "rate_period": ["rate period", "period", "pay period"],
    "category": ["category"],
    "job_type": ["job type", "type", "employment type"],
    "work_arrangement": ["work arrangement", "arrangement", "workplace", "remote"],
    "location": ["location"],
    "eligible_countries": ["eligible countries", "countries", "eligibility"],
    "posted_date": ["posted date", "posted", "date posted", "posted on", "published", "published at"],
    "expiry_date": ["expiry date", "expires", "expiry", "expires on", "deadline", "closing date"],
    "status": ["status"],
    "last_verified_at": ["last verified at", "last verified", "verified at", "verified on", "verified"],
    "referral_link": ["referral link", "referral url", "referral"],
}
REQUIRED = ("title", "platform", "job_link")
TEMPLATE_HEADERS = list(COLUMNS)
TEMPLATE_EXAMPLE = {
    "external_id": "EX-1001",
    "title": "Example: AI Trainer — Python",
    "platform": "Example Company",
    "job_link": "https://example.com/jobs/123",
    "job_description": "Delete this example row before uploading. Review and rate AI-written Python code.",
    "hourly_rate": "",
    "minimum_rate": "40",
    "maximum_rate": "60",
    "currency": "USD",
    "rate_period": "hour",
    "category": "AI coding & software evaluation",
    "job_type": "contract",
    "work_arrangement": "remote",
    "location": "Worldwide",
    "eligible_countries": "",
    "posted_date": "2026-10-01",
    "expiry_date": "2026-12-31",
    "status": "open",
    "last_verified_at": "2026-10-27",
    "referral_link": "",
}
# Shown in the template's "How to fill" sheet and on the admin upload page.
COLUMN_HELP = [
    ("title", "yes", "Job title as shown on the original listing (max 200 characters)."),
    (
        "platform",
        "yes",
        "Company / platform, e.g. Mercor, Micro1, Turing, DataAnnotation. Matched to existing names ignoring case and spaces; new names are added.",
    ),
    ("job_link", "yes", "Full https:// link to the original listing. Used to recognise the same job in later uploads."),
    ("external_id", "no", "The platform's own job ID, if you have one. When present it is used (with the platform) to match re-uploads."),
    ("job_description", "no", "Plain text or simple HTML."),
    ("hourly_rate", "no", "A number (40 = 40 per hour) or text such as $40/hr, $40–60/hr, Up to $100/hr."),
    ("minimum_rate / maximum_rate", "no", "Numbers. Use with currency and rate_period for a range."),
    ("currency", "no", "ISO code (USD, GBP, EUR, INR, CAD, AUD, PKR…) or symbol. Defaults to USD when a rate is given."),
    ("rate_period", "no", "hour, task, week, month, year or project. Defaults to hour."),
    ("category", "no", "An existing category name. If empty or unknown, one is suggested from the title."),
    ("job_type", "no", "contract, freelance, hourly, part-time, full-time or internship."),
    ("work_arrangement", "no", "remote, hybrid or onsite. Defaults to remote."),
    ("location", "no", "Free text, e.g. Worldwide, or Remote — US & Canada."),
    ("eligible_countries", "no", "Where applicants must live: Worldwide, or country names / ISO codes separated by commas (US, GB, Pakistan)."),
    (
        "posted_date / expiry_date / last_verified_at",
        "no",
        "Dates: YYYY-MM-DD (also DD/MM/YYYY or real Excel dates). A listing is hidden after its expiry date.",
    ),
    ("status", "no", "open (default for new jobs), closed, draft or expired."),
    ("referral_link", "no", "Your referral URL. Used only while a referral program is active for the platform."),
]

JOB_TYPES = {
    "contract": Job.EmploymentType.CONTRACT,
    "contractor": Job.EmploymentType.CONTRACT,
    "freelance": Job.EmploymentType.FREELANCE,
    "hourly": Job.EmploymentType.HOURLY,
    "task": Job.EmploymentType.HOURLY,
    "task-based": Job.EmploymentType.HOURLY,
    "part-time": Job.EmploymentType.PART_TIME,
    "part time": Job.EmploymentType.PART_TIME,
    "part_time": Job.EmploymentType.PART_TIME,
    "full-time": Job.EmploymentType.FULL_TIME,
    "full time": Job.EmploymentType.FULL_TIME,
    "full_time": Job.EmploymentType.FULL_TIME,
    "internship": Job.EmploymentType.INTERNSHIP,
}
ARRANGEMENTS = {
    "remote": Job.WorkArrangement.REMOTE,
    "hybrid": Job.WorkArrangement.HYBRID,
    "onsite": Job.WorkArrangement.ONSITE,
    "on-site": Job.WorkArrangement.ONSITE,
    "on site": Job.WorkArrangement.ONSITE,
    "office": Job.WorkArrangement.ONSITE,
}
STATUSES = {
    "open": Job.Status.PUBLISHED,
    "active": Job.Status.PUBLISHED,
    "published": Job.Status.PUBLISHED,
    "live": Job.Status.PUBLISHED,
    "closed": Job.Status.CLOSED,
    "filled": Job.Status.CLOSED,
    "inactive": Job.Status.CLOSED,
    "draft": Job.Status.DRAFT,
    "hidden": Job.Status.DRAFT,
    "expired": Job.Status.EXPIRED,
}
YES = {"yes", "y", "true", "1", "verified"}


class SpreadsheetError(Exception):
    """Problem with the file as a whole. Message is safe to show."""


class ImportBlocked(Exception):
    """The import can't be applied as previewed. Message is safe to show."""


# --- Reading --------------------------------------------------------------------------------------


def _key(header):
    return re.sub(r"[\s_\-]+", " ", str(header or "").strip().lower())


def _header_map(headers):
    lookup = {_key(alias): canonical for canonical, aliases in COLUMNS.items() for alias in aliases}
    lookup.update({_key(c): c for c in COLUMNS})
    mapping = {}
    for idx, h in enumerate(headers):
        canonical = lookup.get(_key(h))
        if canonical and canonical not in mapping.values():
            mapping[idx] = canonical
    missing = [c for c in REQUIRED if c not in mapping.values()]
    if missing:
        raise SpreadsheetError(f"Missing required column(s): {', '.join(missing)}. Download the template to see the expected headers.")
    return mapping


def _cell(value):
    if value is None:
        return ""
    if isinstance(value, datetime | date):
        return value.isoformat()[:10]
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def read_rows(uploaded):
    """Return a list of {column: str} dicts, each with `_row` (the spreadsheet row number)."""
    name = (getattr(uploaded, "name", "") or "").lower()
    data = uploaded.read(MAX_BYTES + 1)
    if len(data) > MAX_BYTES:
        raise SpreadsheetError(f"The file is larger than {MAX_BYTES // (1024 * 1024)} MB. Split it into smaller files.")
    if name.endswith(".xlsx") or data[:4] == b"PK\x03\x04":
        raw_rows = _xlsx_rows(data)
    elif name.endswith(".xls"):
        raise SpreadsheetError("Old .xls files aren't supported. In Excel choose File → Save As → .xlsx (or CSV) and upload again.")
    elif name.endswith((".csv", ".txt")) or not name:
        raw_rows = _csv_rows(data)
    else:
        raise SpreadsheetError("Upload an Excel (.xlsx) or CSV file.")

    raw_rows = iter(raw_rows)
    headers = next(raw_rows, None)
    if not headers or not any(_cell(h) for h in headers):
        raise SpreadsheetError("The file is empty.")
    mapping = _header_map(headers)
    rows = []
    for number, values in enumerate(raw_rows, start=2):
        row = {canonical: _cell(values[idx]) if idx < len(values) else "" for idx, canonical in mapping.items()}
        if not any(row.values()):
            continue  # blank line
        if len(rows) >= MAX_ROWS:
            raise SpreadsheetError(f"The file has more than {MAX_ROWS} job rows. Split it into smaller files.")
        if "job_description" in row:
            row["job_description"] = row["job_description"][:20000]
        row["_row"] = number
        rows.append(row)
    if not rows:
        raise SpreadsheetError("No job rows found below the header row.")
    return rows


def _xlsx_rows(data):
    from openpyxl import load_workbook

    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # corrupt / not really xlsx
        raise SpreadsheetError("This doesn't look like a valid .xlsx file.") from exc
    ws = wb.worksheets[0]
    rows = [list(r) for r in ws.iter_rows(values_only=True, max_row=MAX_ROWS + 50)]
    wb.close()
    return rows


def _csv_rows(data):
    text = data.decode("utf-8-sig", errors="replace")
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    return list(csv.reader(io.StringIO(text), dialect))[: MAX_ROWS + 50]


# --- Field parsing --------------------------------------------------------------------------------

CURRENCY_SYMBOLS = {"$": "USD", "£": "GBP", "€": "EUR", "₹": "INR"}
CURRENCY_CODES = {"USD", "GBP", "EUR", "INR", "CAD", "AUD", "PKR", "NZD", "SGD", "AED", "CHF", "JPY"}
PERIODS = [
    (r"/\s*h(?:ou)?r\b|/\s*h\b|per\s+hour|hourly|an\s+hour|p/?h\b", "hour"),
    (r"/\s*task\b|per\s+task", "task"),
    (r"/\s*w(?:ee)?k\b|per\s+week|weekly", "week"),
    (r"/\s*mo(?:nth)?\b|per\s+month|monthly", "month"),
    (r"/\s*y(?:ea)?r\b|per\s+(?:year|annum)|annual|yearly|p\.?a\.?", "year"),
    (r"/\s*project\b|per\s+project", "project"),
]
PERIOD_WORDS = {
    "hour": "hour", "hr": "hour", "h": "hour", "hourly": "hour", "per hour": "hour",
    "task": "task", "per task": "task",
    "week": "week", "wk": "week", "weekly": "week", "per week": "week",
    "month": "month", "mo": "month", "monthly": "month", "per month": "month",
    "year": "year", "yr": "year", "yearly": "year", "annual": "year", "annually": "year", "per year": "year",
    "project": "project", "per project": "project",
}  # fmt: skip
NUM = r"(\d[\d,]*(?:\.\d+)?)\s*(k)?"


def parse_rate(text):
    """'$40–60/hr' → (40, 60, 'USD', 'hour'). Returns None unless amount, currency and period are all clear."""
    t = (text or "").strip()
    if not t:
        return None
    currency = _currency_in(t)
    period = next((p for pattern, p in PERIODS if re.search(pattern, t, re.I)), None)
    nums = []
    for m in re.finditer(NUM, t, re.I):
        try:
            value = Decimal(m.group(1).replace(",", ""))
        except InvalidOperation:
            continue
        nums.append(value * 1000 if m.group(2) else value)
    if not (currency and period and nums) or len(nums) > 2:
        return None
    lo, hi = (nums[0], nums[-1]) if len(nums) == 2 else (nums[0], nums[0])
    if re.search(r"\bup\s+to\b", t, re.I) and len(nums) == 1:
        lo = None
    if lo is not None and lo > hi:
        return None
    return lo, hi, currency, period


def _currency_in(text):
    code = next((code for sym, code in CURRENCY_SYMBOLS.items() if sym in text), None)
    return code or next((c for c in CURRENCY_CODES if re.search(rf"\b{c}\b", text, re.I)), None)


def _amount(text):
    """'40', '$40', '1,200.50' → Decimal; None if it isn't a plain amount."""
    cleaned = re.sub(r"[\s,]", "", re.sub(r"^[^\d.\-]+", "", (text or "").strip()))
    cleaned = re.sub(r"(?i)[a-z/]+$", "", cleaned)
    try:
        value = Decimal(cleaned)
    except InvalidOperation:
        return None
    return value if value.is_finite() else None


def _date(value):
    if not value:
        return None
    d = parse_date(value[:10]) if re.match(r"^\d{4}-\d{2}-\d{2}", value) else None
    if d is None:
        for fmt in ("%d/%m/%Y", "%m/%d/%Y", "%d-%m-%Y", "%d.%m.%Y", "%d %b %Y", "%d %B %Y", "%b %d %Y", "%B %d %Y", "%b %d, %Y", "%B %d, %Y"):
            try:
                return datetime.strptime(value.strip(), fmt).date()
            except ValueError:
                continue
    return d


def _aware(d, at=time(12, 0)):
    return timezone.make_aware(datetime.combine(d, at)) if d else None


def _countries(text):
    """'US, United Kingdom; PK' → (['GB', 'PK', 'US'], unknown_tokens). 'Worldwide' → ('worldwide', [])."""
    if any(w in text.lower() for w in WORLDWIDE_WORDS):
        return "worldwide", []
    codes, unknown = set(), []
    for token in re.split(r"[,;/|\n]+", text):
        t = token.strip()
        if not t:
            continue
        code = t.upper() if t.upper() in COUNTRIES else ("GB" if t.upper() == "UK" else COUNTRY_NAME_TO_CODE.get(t.lower()))
        if code:
            codes.add(code)
        else:
            unknown.append(t)
    return sorted(codes), unknown


def _squash(name):
    return re.sub(r"[^a-z0-9]", "", (name or "").lower())


def normalise_platform_name(name):
    return re.sub(r"\s+", " ", (name or "").strip())[:120]


def find_platform(name, platforms=None):
    """Match 'Data Annotation', 'dataannotation' and 'DataAnnotation' to the same platform."""
    key = _squash(name)
    if not key:
        return None
    if platforms is None:
        platforms = {_squash(p.name): p for p in Platform.objects.only("id", "name", "slug")}
    return platforms.get(key)


def platform_for(name, link):
    """Existing platform for this name, or a new one (website = the link's origin)."""
    platform = find_platform(name)
    if platform:
        return platform
    from urllib.parse import urlparse

    name = normalise_platform_name(name)
    parsed = urlparse(link)
    base = slugify(name)[:120] or "company"
    slug, n = base, 2
    while Platform.objects.filter(slug=slug).exists():
        slug, n = f"{base}-{n}", n + 1
    return Platform.objects.create(name=name, slug=slug, website=f"{parsed.scheme}://{parsed.netloc}")


# --- Validation & normalisation -------------------------------------------------------------------


def _clean(d, errors, warnings, categories):
    """Validate one raw row. Returns the job field values that the row sets (only filled-in cells)."""
    v = {}
    title = d.get("title", "")
    if title.lower().startswith("example:"):
        errors.append("This is the example row from the template — delete it.")
    elif not title:
        errors.append("title is empty.")
    elif len(title) > 200:
        errors.append("title is longer than 200 characters.")
    else:
        v["title"] = re.sub(r"\s+", " ", title)

    if not d.get("platform"):
        errors.append("platform is empty.")
    elif len(d["platform"]) > 120:
        errors.append("platform is longer than 120 characters.")

    link = safe_url(d.get("job_link", ""))
    if not link:
        errors.append(f"job_link '{d.get('job_link', '')[:80]}' must be a full http:// or https:// URL.")
    else:
        v["application_url"] = link
    if d.get("referral_link"):
        if safe_url(d["referral_link"]):
            v["referral_url"] = safe_url(d["referral_link"])
        else:
            errors.append("referral_link must be a full http:// or https:// URL.")

    if d.get("external_id"):
        if len(d["external_id"]) > 200:
            errors.append("external_id is longer than 200 characters.")
        else:
            v["external_id"] = d["external_id"]

    if d.get("job_description"):
        v["description"] = html_to_text(d["job_description"])

    _clean_rate(d, v, errors, warnings)

    if d.get("category"):
        cat = categories.get(d["category"].strip().lower()) or categories.get(slugify(d["category"]))
        if cat:
            v["category"] = cat
        else:
            warnings.append(f"Category '{d['category']}' doesn't exist; one will be suggested from the title.")

    if d.get("job_type"):
        if d["job_type"].lower() in JOB_TYPES:
            v["employment_type"] = JOB_TYPES[d["job_type"].lower()]
        else:
            errors.append(f"job_type '{d['job_type']}' isn't recognised — use contract, freelance, hourly, part-time, full-time or internship.")

    if d.get("location"):
        _, v["eligibility_scope"], v["eligible_countries"] = infer_location(d["location"])
        lowered = d["location"].lower()
        if "hybrid" in lowered:
            v["work_arrangement"] = Job.WorkArrangement.HYBRID
        elif re.search(r"on-?site|in[- ]office", lowered):
            v["work_arrangement"] = Job.WorkArrangement.ONSITE
        else:
            v["work_arrangement"] = Job.WorkArrangement.REMOTE  # the directory is for remote AI-training work
    if d.get("work_arrangement"):
        key = d["work_arrangement"].strip().lower()
        if key in ARRANGEMENTS:
            v["work_arrangement"] = ARRANGEMENTS[key]
        else:
            errors.append(f"work_arrangement '{d['work_arrangement']}' isn't recognised — use remote, hybrid or onsite.")
    if d.get("eligible_countries"):
        codes, unknown = _countries(d["eligible_countries"])
        if unknown:
            errors.append(f"eligible_countries: couldn't recognise {', '.join(unknown[:5])}. Use country names or ISO codes (US, GB, PK).")
        elif codes == "worldwide":
            v["eligibility_scope"], v["eligible_countries"] = Job.Eligibility.WORLDWIDE, []
        else:
            v["eligibility_scope"], v["eligible_countries"] = Job.Eligibility.RESTRICTED, codes

    for column, model_field in (("posted_date", "published_at"), ("expiry_date", "expires_at"), ("last_verified_at", "last_verified_at")):
        raw = d.get(column, "")
        if not raw:
            continue
        if column == "last_verified_at" and raw.lower() in YES:
            v[model_field] = timezone.now()
            continue
        parsed = _date(raw)
        if parsed is None:
            errors.append(f"{column} '{raw}' isn't a date — use YYYY-MM-DD.")
        elif column == "expiry_date":
            v[model_field] = _aware(parsed, time(23, 59))
        elif parsed > timezone.localdate():
            errors.append(f"{column} {parsed} is in the future.")
        else:
            v[model_field] = _aware(parsed)

    if d.get("status"):
        key = d["status"].strip().lower()
        if key in STATUSES:
            v["status"] = STATUSES[key]
        else:
            errors.append(f"status '{d['status']}' isn't recognised — use open, closed, draft or expired.")
    if v.get("expires_at") and v["expires_at"] <= timezone.now() and v.get("status", Job.Status.PUBLISHED) == Job.Status.PUBLISHED:
        warnings.append("expiry_date has passed, so this listing won't be shown.")
    return v


def _clean_rate(d, v, errors, warnings):
    if not any(d.get(c) for c in ("hourly_rate", "minimum_rate", "maximum_rate")):
        return
    text = d.get("hourly_rate", "")
    lo = hi = None
    currency = period = None
    if text:
        amount = _amount(text)
        parsed = None if amount is not None else parse_rate(text)
        if amount is not None:
            lo = hi = amount
            period = "hour"
        elif parsed:
            lo, hi, currency, period = parsed
        else:
            v["rate_text"] = text[:120]
            warnings.append("hourly_rate will be shown exactly as written (couldn't read an amount, currency and period for pay sorting).")
            return
        v["rate_text"] = "" if amount is not None else text[:120]
    for column in ("minimum_rate", "maximum_rate"):
        if d.get(column):
            amount = _amount(d[column])
            if amount is None:
                errors.append(f"{column} '{d[column]}' isn't a number.")
                return
            if column == "minimum_rate":
                lo = amount
            else:
                hi = amount
            currency = currency or _currency_in(d[column])
    if (lo is not None and lo < 0) or (hi is not None and hi < 0):
        errors.append("Rates can't be negative.")
        return
    if lo is not None and hi is not None and lo > hi:
        errors.append(f"minimum_rate {lo} is higher than maximum_rate {hi}.")
        return
    if d.get("currency"):
        code = CURRENCY_SYMBOLS.get(d["currency"].strip(), d["currency"].strip().upper())
        if not re.fullmatch(r"[A-Z]{3}", code):
            errors.append(f"currency '{d['currency']}' isn't a 3-letter code such as USD.")
            return
        currency = code
    if d.get("rate_period"):
        period = PERIOD_WORDS.get(d["rate_period"].strip().lower().lstrip("/"))
        if not period:
            errors.append(f"rate_period '{d['rate_period']}' isn't recognised — use hour, task, week, month, year or project.")
            return
    if not currency:
        currency = "USD"
        warnings.append("No currency given; USD assumed.")
    v.update(
        compensation_min=lo if lo is not None else hi,
        compensation_max=hi if hi is not None else lo,
        compensation_currency=currency,
        compensation_period=period or "hour",
        compensation_verified=True,
    )
    v.setdefault("rate_text", "")


# --- Analysis (dry run) ---------------------------------------------------------------------------


@dataclass
class RowResult:
    number: int
    data: dict
    status: str = "new"  # new | update | unchanged | duplicate | error
    errors: list = field(default_factory=list)
    warnings: list = field(default_factory=list)
    values: dict = field(default_factory=dict)
    existing: Job | None = None
    platform: Platform | None = None
    changes: list = field(default_factory=list)
    note: str = ""

    @property
    def title(self):
        return self.data.get("title", "")


# Model fields compared to decide whether a matched job actually changes.
COMPARED = [
    "title", "application_url", "referral_url", "external_id", "description", "rate_text", "compensation_min", "compensation_max",
    "compensation_currency", "compensation_period", "compensation_verified", "category", "employment_type", "work_arrangement",
    "eligibility_scope", "eligible_countries", "published_at", "expires_at", "status", "last_verified_at",
]  # fmt: skip


def _changes(job, values):
    changed = []
    for name in COMPARED:
        if name not in values:
            continue
        new, old = values[name], getattr(job, name)
        if name in ("published_at", "last_verified_at", "expires_at"):
            same = (old and new and timezone.localdate(old) == timezone.localdate(new)) or old == new
        elif name == "eligible_countries":
            same = sorted(old or []) == sorted(new or [])
        else:
            same = old == new
        if not same:
            changed.append(name)
    return changed


def analyse(rows, mode=ImportRun.Mode.UPSERT):
    """Dry run. Returns (results, jobs_to_close). Writes nothing."""
    platforms = {_squash(p.name): p for p in Platform.objects.only("id", "name", "slug")}
    categories = {}
    for c in Category.objects.all():
        categories[c.name.lower()] = categories[c.slug] = c
    results = []
    for raw in rows:
        r = RowResult(number=raw.get("_row", 0), data=dict(raw))
        r.values = _clean(raw, r.errors, r.warnings, categories)
        r.platform = find_platform(raw.get("platform", ""), platforms)
        if raw.get("platform") and r.platform is None and not r.errors:
            r.warnings.append(f"New platform '{normalise_platform_name(raw['platform'])}' will be added.")
        results.append(r)

    # Load every job these rows could match, in two queries.
    valid = [r for r in results if not r.errors]
    links = {canonical_url(r.values["application_url"]) for r in valid}
    by_link = {}
    for job in Job.objects.filter(canonical_url__in=links).exclude(status=Job.Status.REJECTED).select_related("platform", "category"):
        by_link.setdefault(job.canonical_url, []).append(job)
    ext_ids = {r.values["external_id"] for r in valid if r.values.get("external_id")}
    by_ext = {}
    for job in Job.objects.filter(external_id__in=ext_ids).exclude(status=Job.Status.REJECTED).select_related("platform", "category"):
        by_ext.setdefault((job.platform_id, job.external_id), []).append(job)

    seen = {}
    for r in valid:
        link = canonical_url(r.values["application_url"])
        pkey = r.platform.pk if r.platform else _squash(r.data["platform"])
        keys = [("link", pkey, link)] + ([("ext", pkey, r.values["external_id"])] if r.values.get("external_id") else [])
        earlier = next((seen[k] for k in keys if k in seen), None)
        if earlier:
            r.status, r.note = "duplicate", f"Same job as row {earlier} in this file — skipped."
            continue
        for k in keys:
            seen[k] = r.number

        same_link = by_link.get(link, [])
        candidates = []
        if r.platform and r.values.get("external_id"):
            candidates = by_ext.get((r.platform.pk, r.values["external_id"]), [])
        if not candidates and r.platform:
            candidates = [j for j in same_link if j.platform_id == r.platform.pk]
        if len(candidates) > 1:
            r.errors.append(f"Matches more than one existing job ({', '.join(f'#{j.pk}' for j in candidates)}). Merge them in the admin first.")
            continue
        if not candidates:
            other = [j for j in same_link if not r.platform or j.platform_id != r.platform.pk]
            if other:
                r.errors.append(f"This job_link already belongs to job #{other[0].pk} on {other[0].platform.name}. Check the platform name.")
                continue
            r.status = "new"
            continue
        job = candidates[0]
        clash = [j for j in same_link if j.platform_id == job.platform_id and j.pk != job.pk]
        if clash:
            r.errors.append(f"external_id matches job #{job.pk}, but job_link belongs to job #{clash[0].pk}.")
            continue
        r.existing = job
        r.changes = _changes(job, r.values)
        r.status = "update" if r.changes else "unchanged"

    for r in results:
        if r.errors:
            r.status = "error"

    to_close = []
    if mode == ImportRun.Mode.SNAPSHOT:
        matched = {r.existing.pk for r in results if r.existing is not None}
        to_close = list(snapshot_scope().exclude(pk__in=matched).select_related("platform").order_by("platform__name", "title"))
    return results, to_close


def snapshot_scope():
    """Jobs a full snapshot reconciles: open, non-demo listings managed by uploads/admin (automated feeds are excluded)."""
    return Job.objects.filter(status=Job.Status.PUBLISHED, is_demo=False, source__isnull=True)


def summarize(results, to_close=()):
    count = lambda status: sum(r.status == status for r in results)  # noqa: E731
    return {
        "total": len(results),
        "added": count("new"),
        "updated": count("update"),
        "unchanged": count("unchanged"),
        "duplicates": count("duplicate"),
        "skipped": count("unchanged") + count("duplicate"),
        "rejected": count("error"),
        "closing": len(to_close),
    }


def error_report(results):
    return [{"row": r.number, "title": r.title[:120], "errors": r.errors} for r in results if r.errors]


# --- Writing --------------------------------------------------------------------------------------


def start(rows, *, filename="", mode=ImportRun.Mode.UPSERT, user=None):
    """Store parsed rows on a pending ImportRun (the preview reads them from here)."""
    return ImportRun.objects.create(source=None, filename=filename[:200], mode=mode, rows=rows, status=ImportRun.Status.PENDING, triggered_by=user)


def _write(r):
    job = r.existing or Job(platform=platform_for(r.data["platform"], r.values["application_url"]), status=Job.Status.PUBLISHED)
    for name, value in r.values.items():
        setattr(job, name, value)
    if "work_arrangement" not in r.values and job.pk is None:
        job.work_arrangement = Job.WorkArrangement.REMOTE  # the directory is for remote AI-training work
    if "status" in r.values and r.values["status"] == Job.Status.PUBLISHED:
        job.link_broken = False
        if job.expires_at and job.expires_at <= timezone.now() and "expires_at" not in r.values:
            job.expires_at = None  # re-opened by the sheet; an old expiry would keep it hidden
    if job.category_id is None:
        job.category = infer_category(job.title, job.description)
    if not job.source_url:
        job.source_url = job.application_url
    job.needs_review = False
    job.save()
    return job


def apply(run, *, confirm_close=None):
    """
    Apply a pending ImportRun in a single transaction. Returns the summary dict.

    In snapshot mode `confirm_close` must equal the number of jobs the preview said would be closed; if the data
    changed since the preview, or any row was rejected, nothing is applied.
    """
    if run.status != ImportRun.Status.PENDING:
        raise ImportBlocked("This import has already been applied or cancelled. Upload the file again.")
    results, to_close = analyse(run.rows, run.mode)
    summary = summarize(results, to_close)
    if run.mode == ImportRun.Mode.SNAPSHOT:
        if summary["rejected"]:
            raise ImportBlocked("Full snapshot imports can't run while rows are rejected — fix them (or use “Add & update”) first.")
        if confirm_close != len(to_close):
            raise ImportBlocked(
                f"Confirm closing the {len(to_close)} listing(s) missing from the file. The numbers changed since the preview; review it again."
            )
    run.status = ImportRun.Status.RUNNING
    run.save(update_fields=["status"])
    try:
        with transaction.atomic():
            for r in results:
                if r.status in ("new", "update"):
                    _write(r)
            for job in to_close:
                job.status = Job.Status.CLOSED
                job.save()
    except Exception as exc:
        run.status, run.finished_at = ImportRun.Status.FAILED, timezone.now()
        run.errors = [{"row": None, "title": "", "errors": [f"Import rolled back, nothing was changed: {type(exc).__name__}: {str(exc)[:300]}"]}]
        run.errors += error_report(results)
        run.save()
        raise ImportBlocked("The import failed and was rolled back — nothing was changed. Details are in the import history.") from exc
    run.created_count, run.updated_count = summary["added"], summary["updated"]
    run.duplicate_count, run.failed_count, run.expired_count = summary["skipped"], summary["rejected"], summary["closing"]
    run.errors = error_report(results)
    run.status = ImportRun.Status.PARTIAL if summary["rejected"] else ImportRun.Status.SUCCESS
    run.finished_at = timezone.now()
    run.save()
    return summary


# --- Templates ------------------------------------------------------------------------------------


def template_xlsx():
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Jobs"
    ws.append(TEMPLATE_HEADERS)
    ws.append([TEMPLATE_EXAMPLE[h] for h in TEMPLATE_HEADERS])
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="EEEEEC")
    for idx, header in enumerate(TEMPLATE_HEADERS, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = 48 if header == "job_description" else max(14, len(header) + 4)
    ws.freeze_panes = "A2"
    help_ws = wb.create_sheet("How to fill")
    help_ws.append(["Column", "Required", "What to enter"])
    for line in COLUMN_HELP:
        help_ws.append(list(line))
    help_ws.append([])
    help_ws.append(["", "", "Header names are flexible: e.g. 'Company' works for platform and 'URL' for job_link. Other columns are ignored."])
    help_ws.append(["", "", "Re-uploading the same file is safe: matching jobs are updated in place, never duplicated."])
    help_ws.column_dimensions["A"].width, help_ws.column_dimensions["B"].width, help_ws.column_dimensions["C"].width = 40, 10, 120
    for cell in help_ws[1]:
        cell.font = Font(bold=True)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def template_csv():
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(TEMPLATE_HEADERS)
    writer.writerow([TEMPLATE_EXAMPLE[h] for h in TEMPLATE_HEADERS])
    return buf.getvalue().encode("utf-8-sig")
