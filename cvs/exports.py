"""ATS-friendly DOCX and PDF rendering: single column, standard headings, real text, no images or tables."""

import io
import os
from xml.sax.saxutils import escape

from .schemas import RenderedCV

SECTION_TITLES = {
    "summary": "Professional Summary",
    "skills": "Skills",
    "experience": "Experience",
    "projects": "Projects",
    "education": "Education",
    "certifications": "Certifications",
    "languages": "Languages",
    "other": None,
}


def _dates(start, end):
    if start and end:
        return f"{start} – {end}"
    return start or end or ""


def _contact_line(cv: RenderedCV):
    parts = [cv.contact.email, cv.contact.phone, cv.contact.location, *cv.contact.links]
    return "  |  ".join(p for p in parts if p)


def _has(cv, section):
    return {
        "summary": bool(cv.summary),
        "skills": bool(cv.skills),
        "experience": bool(cv.experience),
        "projects": bool(cv.projects),
        "education": bool(cv.education),
        "certifications": bool(cv.certifications),
        "languages": bool(cv.languages),
        "other": bool(cv.other_sections),
    }[section]


# --- DOCX -------------------------------------------------------------------------------------


def render_docx(cv: RenderedCV) -> bytes:
    import docx
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, RGBColor

    doc = docx.Document()
    for section in doc.sections:
        section.left_margin = section.right_margin = docx.shared.Inches(0.8)
        section.top_margin = section.bottom_margin = docx.shared.Inches(0.7)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(2)

    def heading(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text.upper())
        run.bold = True
        run.font.size = Pt(11.5)
        run.font.color.rgb = RGBColor(0x1F, 0x2A, 0x44)

    def bullet(text):
        p = doc.add_paragraph(text, style="List Bullet")
        p.paragraph_format.space_after = Pt(1)

    def entry_header(left, right, sub=""):
        p = doc.add_paragraph()
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.space_before = Pt(4)
        r = p.add_run(left)
        r.bold = True
        if right:
            p.add_run(f"  ·  {right}").italic = True
        if sub:
            s = doc.add_paragraph(sub)
            s.paragraph_format.keep_with_next = True

    name = doc.add_paragraph()
    name.alignment = WD_ALIGN_PARAGRAPH.LEFT
    nr = name.add_run(cv.full_name or "Curriculum Vitae")
    nr.bold = True
    nr.font.size = Pt(20)
    if cv.headline:
        doc.add_paragraph(cv.headline).runs[0].font.size = Pt(11.5)
    contact = _contact_line(cv)
    if contact:
        doc.add_paragraph(contact).runs[0].font.size = Pt(9.5)

    for section in cv.section_order:
        if section not in SECTION_TITLES or not _has(cv, section):
            continue
        if section == "summary":
            heading(SECTION_TITLES[section])
            doc.add_paragraph(cv.summary)
        elif section == "skills":
            heading("Skills")
            doc.add_paragraph(", ".join(cv.skills))
        elif section == "experience":
            heading("Experience")
            for e in cv.experience:
                title = " — ".join(x for x in [e.title, e.employer] if x)
                entry_header(title, _dates(e.start_date, e.end_date), e.location)
                for b in e.bullets:
                    bullet(b)
        elif section == "projects":
            heading("Projects")
            for p in cv.projects:
                entry_header(p.name, ", ".join(p.technologies))
                if p.description:
                    doc.add_paragraph(p.description)
                for b in p.bullets:
                    bullet(b)
        elif section == "education":
            heading("Education")
            for e in cv.education:
                qual = ", ".join(x for x in [e.qualification, e.field] if x)
                entry_header(" — ".join(x for x in [qual, e.institution] if x), _dates(e.start_date, e.end_date))
                for d in e.details:
                    bullet(d)
        elif section == "certifications":
            heading("Certifications")
            for c in cv.certifications:
                bullet(" — ".join(x for x in [c.name, c.issuer, c.date] if x))
        elif section == "languages":
            heading("Languages")
            doc.add_paragraph(", ".join(cv.languages))
        elif section == "other":
            for s in cv.other_sections:
                heading(s.title)
                for item in s.items:
                    bullet(item)

    buf = io.BytesIO()
    doc.core_properties.title = f"CV — {cv.full_name}" if cv.full_name else "CV"
    doc.save(buf)
    return buf.getvalue()


# --- PDF ----------------------------------------------------------------------------------------

_FONTS_REGISTERED = False


def _register_fonts():
    """Use the Bitstream Vera TTFs bundled with ReportLab (broader glyph coverage than Helvetica)."""
    global _FONTS_REGISTERED
    if _FONTS_REGISTERED:
        return "CVSans", "CVSans-Bold", "CVSans-Italic"
    import reportlab
    from reportlab.lib.fonts import addMapping
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    font_dir = os.path.join(os.path.dirname(reportlab.__file__), "fonts")
    try:
        pdfmetrics.registerFont(TTFont("CVSans", os.path.join(font_dir, "Vera.ttf")))
        pdfmetrics.registerFont(TTFont("CVSans-Bold", os.path.join(font_dir, "VeraBd.ttf")))
        pdfmetrics.registerFont(TTFont("CVSans-Italic", os.path.join(font_dir, "VeraIt.ttf")))
        pdfmetrics.registerFont(TTFont("CVSans-BoldItalic", os.path.join(font_dir, "VeraBI.ttf")))
        addMapping("CVSans", 0, 0, "CVSans")
        addMapping("CVSans", 1, 0, "CVSans-Bold")
        addMapping("CVSans", 0, 1, "CVSans-Italic")
        addMapping("CVSans", 1, 1, "CVSans-BoldItalic")
        _FONTS_REGISTERED = True
        return "CVSans", "CVSans-Bold", "CVSans-Italic"
    except Exception:
        return "Helvetica", "Helvetica-Bold", "Helvetica-Oblique"


def render_pdf(cv: RenderedCV) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import KeepTogether, ListFlowable, ListItem, Paragraph, SimpleDocTemplate, Spacer

    regular, bold, italic = _register_fonts()
    navy = colors.HexColor("#1F2A44")
    styles = {
        "name": ParagraphStyle("name", fontName=bold, fontSize=19, leading=23, textColor=navy),
        "headline": ParagraphStyle("headline", fontName=regular, fontSize=11, leading=14, spaceBefore=2),
        "contact": ParagraphStyle("contact", fontName=regular, fontSize=8.8, leading=11, textColor=colors.HexColor("#444444"), spaceBefore=3),
        "h": ParagraphStyle("h", fontName=bold, fontSize=11, leading=14, textColor=navy, spaceBefore=10, spaceAfter=3),
        "entry": ParagraphStyle("entry", fontName=bold, fontSize=10, leading=13, spaceBefore=5),
        "sub": ParagraphStyle("sub", fontName=italic, fontSize=9, leading=12, textColor=colors.HexColor("#555555")),
        "body": ParagraphStyle("body", fontName=regular, fontSize=9.6, leading=12.6, alignment=TA_LEFT),
        "bullet": ParagraphStyle("bullet", fontName=regular, fontSize=9.6, leading=12.4),
    }
    esc = lambda s: escape(s or "")  # noqa: E731

    story = [Paragraph(esc(cv.full_name or "Curriculum Vitae"), styles["name"])]
    if cv.headline:
        story.append(Paragraph(esc(cv.headline), styles["headline"]))
    contact = _contact_line(cv)
    if contact:
        story.append(Paragraph(esc(contact), styles["contact"]))

    def bullets(items):
        return ListFlowable(
            [ListItem(Paragraph(esc(b), styles["bullet"]), leftIndent=10) for b in items],
            bulletType="bullet",
            start="•",
            leftIndent=10,
            bulletFontName=regular,
            bulletFontSize=8,
        )

    def entry(title, right="", sub="", items=(), body=""):
        block = [Paragraph(esc(title) + (f'<font name="{italic}" size="9">  ·  {esc(right)}</font>' if right else ""), styles["entry"])]
        if sub:
            block.append(Paragraph(esc(sub), styles["sub"]))
        if body:
            block.append(Paragraph(esc(body), styles["body"]))
        head = KeepTogether(block)
        return [head, bullets(items)] if items else [head]

    for section in cv.section_order:
        if section not in SECTION_TITLES or not _has(cv, section):
            continue
        if section != "other":
            story.append(Paragraph(SECTION_TITLES[section].upper(), styles["h"]))
        if section == "summary":
            story.append(Paragraph(esc(cv.summary), styles["body"]))
        elif section == "skills":
            story.append(Paragraph(esc(", ".join(cv.skills)), styles["body"]))
        elif section == "experience":
            for e in cv.experience:
                story += entry(" — ".join(x for x in [e.title, e.employer] if x), _dates(e.start_date, e.end_date), e.location, e.bullets)
        elif section == "projects":
            for p in cv.projects:
                story += entry(p.name, ", ".join(p.technologies), "", p.bullets, p.description)
        elif section == "education":
            for e in cv.education:
                qual = ", ".join(x for x in [e.qualification, e.field] if x)
                story += entry(" — ".join(x for x in [qual, e.institution] if x), _dates(e.start_date, e.end_date), "", e.details)
        elif section == "certifications":
            story.append(bullets([" — ".join(x for x in [c.name, c.issuer, c.date] if x) for c in cv.certifications]))
        elif section == "languages":
            story.append(Paragraph(esc(", ".join(cv.languages)), styles["body"]))
        elif section == "other":
            for s in cv.other_sections:
                story.append(Paragraph(esc(s.title.upper()), styles["h"]))
                story.append(bullets(s.items))
        story.append(Spacer(1, 1 * mm))

    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf,
        pagesize=A4,
        leftMargin=18 * mm,
        rightMargin=18 * mm,
        topMargin=16 * mm,
        bottomMargin=16 * mm,
        title=f"CV — {cv.full_name}" if cv.full_name else "CV",
        author=cv.full_name or "",
    )
    doc.build(story)
    return buf.getvalue()
