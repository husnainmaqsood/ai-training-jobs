import io
import zipfile

import docx
import pytest
from pypdf import PdfReader

from cvs.exports import render_docx, render_pdf
from cvs.parsing import heuristic_structure
from cvs.schemas import DEFAULT_SECTION_ORDER, RenderedCV


@pytest.fixture
def rendered(sample_cv_text):
    s = heuristic_structure(sample_cv_text)
    return RenderedCV(**s.model_dump(exclude={"headline"}), headline="Data specialist <script>&", section_order=DEFAULT_SECTION_ORDER)


def test_pdf_has_selectable_text(rendered):
    data = render_pdf(rendered)
    assert data.startswith(b"%PDF")
    text = "\n".join(p.extract_text() for p in PdfReader(io.BytesIO(data)).pages)
    assert "Example Analytics Ltd" in text and "EXPERIENCE" in text and "<script>&" in text


def test_docx_is_valid_and_structured(rendered):
    data = render_docx(rendered)
    assert zipfile.is_zipfile(io.BytesIO(data))
    d = docx.Document(io.BytesIO(data))
    texts = [p.text for p in d.paragraphs]
    assert "EXPERIENCE" in texts and any("Example Analytics Ltd" in t for t in texts)
    assert not d.inline_shapes and not d.tables  # ATS-friendly: no images or tables


def test_long_cv_paginates(rendered):
    rendered.experience = rendered.experience * 15
    reader = PdfReader(io.BytesIO(render_pdf(rendered)))
    assert len(reader.pages) > 1


def test_unicode_names(rendered):
    rendered.full_name = "Zoë Ångström Muñoz"
    text = PdfReader(io.BytesIO(render_pdf(rendered))).pages[0].extract_text()
    assert "Zoë Ångström Muñoz" in text


@pytest.mark.django_db
def test_docx_export_requires_premium(client_user, user, master_cv, subscription):
    from cvs.services import start_generation

    gen = start_generation(user=user, master=master_cv, target_title="T", job_description="Python SQL evaluation work for AI. " * 8)
    gcv = gen.versions.get()
    resp = client_user.get(f"/cv/tailored/{gcv.public_id}/download/docx/")
    assert resp.status_code == 200 and resp["Content-Type"].startswith("application/vnd.openxmlformats")
    subscription.status = "canceled"
    subscription.save()
    resp = client_user.get(f"/cv/tailored/{gcv.public_id}/download/docx/")
    assert resp.status_code == 302
